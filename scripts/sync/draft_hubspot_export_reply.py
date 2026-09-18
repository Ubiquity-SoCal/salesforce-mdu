"""
Park an Outlook reply to Kia Zaman's 9/14 request for the full Salesforce export and the
property / structure details, as a reply in her own thread. Drafts only, never sent.

Every number in the body is read from the export folder and the field map at run time,
so rerunning after a fresh export (at cutover, say) gives a correct email, never a stale one.

The export itself is NOT attached: it is every contact's name, email and phone. The body
carries a highlighted placeholder for the SharePoint link Koa shares with Kia only.

Run: python draft_hubspot_export_reply.py [--link <share url>]
"""
import csv
import sys
from pathlib import Path

sys.path.insert(0, r"C:\Users\cass\Work_Projects")
from _shared.outlook_draft import open_draft  # noqa: E402

csv.field_size_limit(10**9)
SF = Path(__file__).resolve().parents[2]
EXPORTS = SF / "data" / "output" / "hubspot-export"
FIELD_MAP = SF / "data" / "output" / "salesforce-hubspot-field-map.xlsx"
TO = "kzaman@fiberfirst.com"
THREAD_MARKER = "Complete record exports"
SUBJECT = "Salesforce export and field details for the HubSpot migration"

TH = 'style="background:#305496;color:#fff;font-weight:bold;text-align:left;padding:5px 8px;border:1px solid #bfbfbf"'
TD = 'style="padding:5px 8px;border:1px solid #bfbfbf;vertical-align:top"'
TDN = 'style="padding:5px 8px;border:1px solid #bfbfbf;vertical-align:top;text-align:center;font-weight:bold"'
AMBER = "#FFF3CD"
TABLE = 'cellpadding="0" cellspacing="0" style="border-collapse:collapse;font-family:Calibri,Arial;font-size:10.5pt"'


def rows(folder, name):
    with open(folder / name, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def n(v):
    return f"{v:,}"


def table(header, body, fill=None):
    h = "".join(f"<th {TH}>{c}</th>" for c in header)
    out = []
    for r in body:
        style = f' style="background:{fill}"' if fill else ""
        cells = "".join(f"<td {TDN if isinstance(c, int) else TD}>{n(c) if isinstance(c, int) else c}</td>" for c in r)
        out.append(f"<tr{style}>{cells}</tr>")
    return f"<table {TABLE}><tr>{h}</tr>{''.join(out)}</table>"


def build_body(folder, link=None):
    link_html = (f'<a href="{link}">shared folder</a>' if link
                 else 'shared folder <span style="background:#FFFF00"><b>[folder link]</b></span>')
    man = {r["file"]: r for r in rows(folder, "manifest.csv")}
    cnt = lambda f: int(man[f]["rows_exported"])  # noqa: E731
    rel = {(r["file"], r["field"]): r for r in rows(folder, "relationship-check.csv")}

    contacts = rows(folder, "Contact.csv")
    shared = [c for c in contacts if c["EXPORT_SameEmailCount"] and int(c["EXPORT_SameEmailCount"]) > 1]
    shared_emails = len({c["Email"].strip().lower() for c in shared})
    no_email = sum(1 for c in contacts if not c["Email"].strip())

    accounts = rows(folder, "Account.csv")
    shared_dom = [a for a in accounts if a["EXPORT_SameDomainCount"] and int(a["EXPORT_SameDomainCount"]) > 1]
    generic = sum(1 for a in shared_dom if a["EXPORT_DomainIsGeneric"] == "true")

    opps = rows(folder, "Opportunity.csv")
    inactive = sum(1 for o in opps if o["EXPORT_OwnerIsActive"] == "false")
    archived = sum(1 for t in rows(folder, "Task.csv") if t["IsArchived"] == "true")

    ironclad = int(rel[("Agreement__c.csv", "IronClad_Record__c")]["points_outside_export"])
    leads = sum(int(rel[(f, "WhoId")]["points_outside_export"]) for f in ("Task.csv", "Event.csv") if (f, "WhoId") in rel)
    unit_bs = int(rel[("Property_Unit__c.csv", "Opportunity__c")]["points_at_excluded_business_sales"])
    bs_opps = int(man["Opportunity.csv"]["rows_excluded"])

    title_only = sum(1 for r in rows(folder, "Notes.csv") if not r["BodyText"].strip())
    special = int(rel.get(("Contact.csv", "Special_Project__c"), {}).get("points_outside_export", 0))
    stages = rows(folder, "OpportunityStage.csv")
    open_as_closed = sum(1 for s in stages if s["IsActive"] == "true" and s["IsClosed"] == "false"
                         and s["ForecastCategoryName"] == "Closed")

    formulas = rollups = 0
    try:
        from openpyxl import load_workbook
        wb = load_workbook(FIELD_MAP, read_only=True)
        for ws in wb.worksheets:
            for r in ws.iter_rows(min_row=5, max_col=3, values_only=True):
                if isinstance(r[2], str) and r[2].startswith("formula ("):
                    formulas += 1
                elif isinstance(r[2], str) and r[2].startswith("roll-up summary ("):
                    rollups += 1
    except Exception as e:
        raise SystemExit(f"could not read the field map for the formula count: {e}")
    if not rollups:
        raise SystemExit("field map has no roll-up summary rows: rebuild it with build_hubspot_migration_mapping.py first")

    asked = table(["You asked for", "What is in the folder"], [
        ["Complete record exports",
         "One CSV per object, every field, one row per record, with 18-character Salesforce IDs on every "
         "record, parent, owner and relationship. Columns starting <code>EXPORT_</code> are added helpers "
         "(owner and record type names, duplicate flags); map them or ignore them."],
        ["Activities, notes, files and attachments",
         f"<b>{n(cnt('Task.csv'))}</b> tasks, <b>{n(cnt('Event.csv'))}</b> events, <b>{n(cnt('EmailMessage.csv'))}</b> emails "
         f"and <b>{n(cnt('Notes.csv'))}</b> notes with the full text. <code>ContentDocumentLink.csv</code> ties each note "
         "to its record. Files and attachments are left out: the signed agreements are already in IronClad."
         if "Files.csv" not in man else
         f"<b>{n(cnt('Task.csv'))}</b> tasks, <b>{n(cnt('Event.csv'))}</b> events, <b>{n(cnt('EmailMessage.csv'))}</b> emails, "
         f"<b>{n(cnt('Notes.csv'))}</b> notes with the full text, <b>{n(cnt('Files.csv'))}</b> files and "
         f"<b>{n(cnt('Attachment.csv'))}</b> email attachments. <code>ContentDocumentLink.csv</code> ties each note "
         "and file to its record."],
        ["Property and structure details",
         "The field map, updated: formulas in full, defaults, required and unique settings, picklist values, "
         "dependent picklists, stages per record type, and every relationship with whether its target is in the export."],
    ])

    traps = table(["Before you import", "Count", "What happens"], [
        ["Contacts sharing an email address", len(shared),
         f"{n(shared_emails)} addresses. HubSpot matches contacts on email, so each group becomes one contact. "
         "Flagged in <code>EXPORT_SameEmailCount</code>."],
        ["Contacts with no email", no_email, "They import, but HubSpot cannot deduplicate them."],
        ["Accounts sharing a website domain", len(shared_dom),
         f"They merge into one company if Website maps to Company domain name. {n(generic)} of them are gmail, aol, "
         "facebook and similar (<code>EXPORT_DomainIsGeneric</code>)."],
        ["Opportunities owned by deactivated users", inactive,
         "Still an open owner decision on our side. The original owner is in every row either way."],
        ["Notes with a title and no body", title_only,
         "Salesforce notes have a title; HubSpot notes only have a body. Import <code>EXPORT_NoteForHubSpot</code> "
         "(title plus body), not <code>BodyText</code>, or these arrive blank and every other note loses its title."],
        ["Archived tasks", archived, "Included. A standard Salesforce export leaves these out."],
        ["Formula and roll-up fields", formulas + rollups,
         f"{n(formulas)} formulas and {n(rollups)} roll-ups, exported as today's value. They will not recalculate in "
         "HubSpot unless rebuilt; each definition is in the field map."],
        ["Open stages with Forecast Category set to Closed", open_as_closed,
         "Every stage in Salesforce is set to Closed, including Prospects. Use the Closed and Won columns on the "
         "Stages tab to decide open versus closed."],
        ["Links to things not being migrated", ironclad + special + leads,
         f"{n(ironclad)} agreements point at IronClad records (IronClad moves by its own sync), {n(special)} contacts "
         f"at Special Projects and {n(leads)} activities at Leads. The IDs are kept but will not connect."],
    ], fill=AMBER)

    return f"""
<div style="font-family:Calibri,Arial;font-size:11pt">
<p>Hi Kia,</p>

<p>Both sets are ready in the {link_html}: the field map workbook, and the export in the
<b>salesforce-export-{folder.name}</b> subfolder, one CSV per object. I have not attached them because the
export holds every contact's details.
If you open a CSV in Excel, avoid saving it back: Excel strips leading zeros from zip codes.</p>

{asked}
<br>
<p><b>Business Sales is excluded.</b> {n(bs_opps)} opportunities plus their tasks, emails and notes, per our
8/31 call. Business ROE is a different record type and is fully included. {n(unit_bs)} property units
still point at a Business Sales opportunity; they are in the export and that one link will not connect.</p>

{traps}

<p>This is a snapshot as of {folder.name}. The export is scripted, so I will run the same thing again for the
final pull at cutover. Let me know the date you want it.</p>

<p>Let me know what you find.</p>
<p>Thanks,</p>
</div>
"""


def find_thread():
    """Kia's 9/14 request, so the reply sits in her thread. Items.Restrict on received
    time, never astimezone on ReceivedTime (Outlook hands back naive local times)."""
    import win32com.client
    ns = win32com.client.Dispatch("Outlook.Application").GetNamespace("MAPI")
    items = ns.GetDefaultFolder(6).Items
    items.Sort("[ReceivedTime]", True)
    recent = items.Restrict("[ReceivedTime] >= '09/10/2026 12:00 AM'")
    for m in recent:
        try:
            if m.Class == 43 and THREAD_MARKER in (m.Body or "") and "hubspot" in (m.Body or "").lower():
                return m
        except Exception:
            continue
    return None


def clear_own_prior_draft():
    """A rerun replaces this script's draft instead of stacking copies. Only an exact
    subject match addressed to Kia (or blank) goes, so a draft re-addressed by hand stays."""
    import win32com.client
    drafts = win32com.client.Dispatch("Outlook.Application").GetNamespace("MAPI").GetDefaultFolder(16)
    for item in list(drafts.Items):
        try:
            if item.Subject == SUBJECT and (item.To or "").strip().lower() in {"", TO, "kia zaman"}:
                print(f"  replacing prior draft from {item.CreationTime}")
                item.Delete()
        except Exception as exc:
            print(f"  note: could not inspect a draft: {exc}")


def main():
    folders = sorted(p for p in EXPORTS.iterdir() if p.is_dir() and p.name[:2] == "20")
    if not folders:
        raise SystemExit("no export folder yet, run export_hubspot_records.py first")
    folder = folders[-1]
    # --link <url>: the share link Koa made (9/14: SharePoint folder on PMO_Projects).
    link = sys.argv[sys.argv.index("--link") + 1] if "--link" in sys.argv else None
    body = build_body(folder, link)
    for bad in ("\u2014", "\u2013", "&mdash;", "&ndash;", "&#8212;", "&#8211;", "--"):
        if bad in body:
            raise SystemExit(f"body contains a dash Koa does not want: {bad!r}")

    msg = find_thread()
    if msg is None:
        # 9/14: her request was not in any Outlook folder (came by another channel).
        print("  Kia's message not found in the Inbox since 9/10, parking a NEW draft instead")
        clear_own_prior_draft()
        open_draft(subject=SUBJECT, body_html=body, to=TO, display=False, save=True)
    else:
        print(f"  replying in thread: {msg.Subject!r} from {msg.SenderName} at {msg.ReceivedTime}")
        reply = msg.ReplyAll()
        _ = reply.GetInspector  # makes Outlook insert the reply signature and quoted thread
        existing = reply.HTMLBody or ""
        low = existing.lower()
        b = low.find("<body")
        gt = low.find(">", b) if b != -1 else -1
        reply.HTMLBody = existing[:gt + 1] + body + existing[gt + 1:] if gt != -1 else body + existing
        reply.Save()
    print(f"draft saved to Outlook Drafts (not sent), numbers from {folder}")


if __name__ == "__main__":
    main()
